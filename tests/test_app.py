import io

import pytest

from app import create_app
from app.config import Config
from app.extensions import db, socketio
from app.models import Message, User


class TestConfig(Config):
    TESTING = True
    SQLALCHEMY_DATABASE_URI = "sqlite://"
    SQLALCHEMY_ENGINE_OPTIONS = {}
    WTF_CSRF_ENABLED = False
    RATELIMIT_ENABLED = False
    REDIS_URL = ""
    S3_BUCKET = ""
    SECRET_KEY = "test"


@pytest.fixture()
def app(tmp_path):
    TestConfig.UPLOAD_FOLDER = str(tmp_path)
    from app import throttle

    throttle._local.clear()
    throttle._local_hits.clear()  # cooldowns/flood counters are per-process memory; start each test clean
    app = create_app(TestConfig)
    with app.app_context():
        db.create_all()
        admin = User(name="Admin", email="member.admin@gmail.com", role="admin")
        admin.set_password("adminpass12345")
        db.session.add(admin)
        db.session.commit()
    yield app


@pytest.fixture()
def client(app):
    return app.test_client()


def signup(c, email="member.a@gmail.com", pw="Password@12345", name="Asha", dob="1999-05-17", occupation="Engineer"):
    return c.post("/register", data={"name": name, "email": email, "dob": dob, "occupation": occupation,
                                     "password": pw, "confirm": pw},
                  follow_redirects=True)


def login(c, email, pw):
    return c.post("/login", data={"email": email, "password": pw}, follow_redirects=True)


def test_register_login_and_hashing(app, client):
    assert signup(client).status_code == 200
    with app.app_context():
        u = User.query.filter_by(email="member.a@gmail.com").one()
        assert u.password_hash != "Password@12345" and u.check_password("Password@12345")
        assert str(u.dob) == "1999-05-17" and u.occupation == "Engineer"


def test_home_feed_shows_admin_posts_and_pins(app, client):
    from app.models import Post
    with app.app_context():
        admin = User(name="Boss", email="member.boss@gmail.com", role="admin"); admin.set_password("Password@12345")
        db.session.add(admin); db.session.commit()
        db.session.add(Post(title="Old pinned", body="p", category="info", pinned=True, author_id=admin.id))
        db.session.add(Post(title="Fresh news", body="n", category="news", author_id=admin.id))
        db.session.commit()
    html = client.get("/").get_data(as_text=True)  # visible to visitors who are not logged in
    assert "Fresh news" in html and "Old pinned" in html and "Pinned" in html
    assert html.index("Old pinned") < html.index("Fresh news")


def test_only_admin_can_post_to_feed(app, client):
    signup(client)
    assert client.get("/admin/posts/new").status_code in (302, 403)
    assert client.post("/admin/posts/new", data={"title": "x", "body": "y"}).status_code in (302, 403)


def test_register_requires_dob_and_occupation(app, client):
    for bad in ({"dob": ""}, {"dob": "2999-01-01"}, {"dob": "not-a-date"}, {"occupation": "  "}):
        assert signup(client, **bad).status_code == 400, bad
    with app.app_context():
        assert User.query.filter_by(email="member.a@gmail.com").count() == 0


def test_duplicate_and_short_password_rejected(client):
    signup(client)
    assert signup(client.application.test_client()).status_code == 400
    assert signup(client.application.test_client(), email="member.b@gmail.com", pw="short").status_code == 400


def test_pages_require_login(client):
    for url in ("/community", "/quizzes", "/polls", "/opinions", "/admin/"):
        assert client.get(url).status_code in (302, 401)


def test_admin_only(client):
    signup(client)
    assert client.get("/admin/").status_code == 403
    assert client.get("/admin/users").status_code == 403
    assert client.post("/admin/quizzes/new", data={}).status_code == 403


def test_open_redirect_blocked(client):
    r = client.post("/login?next=https://evil.com", data={"email": "member.admin@gmail.com", "password": "adminpass12345"})
    assert "evil.com" not in r.headers["Location"]


def test_suspended_user_cannot_login_and_is_kicked(app, client):
    signup(client)
    admin = app.test_client()
    login(admin, "member.admin@gmail.com", "adminpass12345")
    with app.app_context():
        uid = User.query.filter_by(email="member.a@gmail.com").one().id
    admin.post(f"/admin/users/{uid}/suspend")
    assert client.get("/community").status_code in (302, 401)  # existing session is dead
    assert login(app.test_client(), "member.a@gmail.com", "Password@12345").status_code == 403


def test_admin_reset_password_forces_change(app, client):
    signup(client)
    admin = app.test_client()
    login(admin, "member.admin@gmail.com", "adminpass12345")
    with app.app_context():
        uid = User.query.filter_by(email="member.a@gmail.com").one().id
    r = admin.post(f"/admin/users/{uid}/reset-password", follow_redirects=True)
    assert b"Temporary password" in r.data
    with app.app_context():
        assert User.query.get(uid).must_change_password


def test_admin_cannot_suspend_self(app):
    admin = app.test_client()
    login(admin, "member.admin@gmail.com", "adminpass12345")
    with app.app_context():
        uid = User.query.filter_by(email="member.admin@gmail.com").one().id
    admin.post(f"/admin/users/{uid}/suspend")
    with app.app_context():
        assert User.query.get(uid).active


def make_quiz(admin):
    import json
    qs = [{"text": "2+2?", "options": ["3", "4", "5"], "correct": 1},
          {"text": "Capital of India?", "options": ["Delhi", "Pune"], "correct": 0}]
    return admin.post("/admin/quizzes/new", data={"title": "Basics", "questions_json": json.dumps(qs)},
                      follow_redirects=True)


def test_quiz_scoring_once_only(app, client):
    admin = app.test_client()
    login(admin, "member.admin@gmail.com", "adminpass12345")
    assert make_quiz(admin).status_code == 200
    signup(client)
    page = client.get("/quizzes/1")
    assert b"2+2?" in page.data
    r = client.post("/quizzes/1", data={"q1": "1", "q2": "1"})
    assert b"1 / 2" in r.data
    r2 = client.post("/quizzes/1", data={"q1": "1", "q2": "0"})  # retake attempt is ignored
    assert b"1 / 2" in r2.data


def test_quiz_validation(app):
    admin = app.test_client()
    login(admin, "member.admin@gmail.com", "adminpass12345")
    r = admin.post("/admin/quizzes/new", data={"title": "x", "questions_json": "[]"})
    assert r.status_code == 400


def test_poll_single_vote(app, client):
    admin = app.test_client()
    login(admin, "member.admin@gmail.com", "adminpass12345")
    admin.post("/admin/polls/new", data={"question": "Bull or bear?", "options": "Bull\nBear"})
    signup(client)
    client.post("/polls/1/vote", data={"option": "1"})
    client.post("/polls/1/vote", data={"option": "2"})  # second vote must not count
    page = client.get("/polls").data
    assert b"1 vote" in page and b"your vote" in page


def test_upload_rules(app, client):
    signup(client)
    ok = client.post("/community/upload", data={"file": (io.BytesIO(b"\x89PNG\r\n\x1a\n"), "pic.png"), "caption": "hi"},
                     content_type="multipart/form-data")
    assert ok.get_json()["ok"] is True
    bad = client.post("/community/upload", data={"file": (io.BytesIO(b"<svg/>"), "x.svg")},
                      content_type="multipart/form-data")
    assert bad.status_code == 400
    exe = client.post("/community/upload", data={"file": (io.BytesIO(b"MZ"), "evil.exe")},
                      content_type="multipart/form-data")
    assert exe.status_code == 400
    hist = client.get("/community/history").get_json()
    assert hist[-1]["file"]["is_image"] and hist[-1]["body"] == "hi"


def test_muted_user_cannot_upload(app, client):
    signup(client)
    with app.app_context():
        User.query.filter_by(email="member.a@gmail.com").one().muted = True
        db.session.commit()
    r = client.post("/community/upload", data={"file": (io.BytesIO(b"x"), "a.txt")}, content_type="multipart/form-data")
    assert r.status_code == 403


def test_socket_chat_flow(app, client):
    signup(client)
    sio = socketio.test_client(app, flask_test_client=client)
    assert sio.is_connected()
    ack = sio.emit("send_message", {"body": "<b>hello</b>"}, callback=True)
    assert ack["ok"]
    got = [m for m in sio.get_received() if m["name"] == "message"]
    first = got[0]["args"][0] if isinstance(got[0]["args"], list) else got[0]["args"]
    assert got and first["body"] == "<b>hello</b>"  # stored raw; the UI escapes via textContent
    assert sio.emit("send_message", {"body": "   "}, callback=True)["ok"] is False


def test_socket_rejects_anonymous(app):
    assert not socketio.test_client(app).is_connected()


def test_admin_delete_message(app, client):
    signup(client)
    sio = socketio.test_client(app, flask_test_client=client)
    sio.emit("send_message", {"body": "spam"}, callback=True)
    admin = app.test_client()
    login(admin, "member.admin@gmail.com", "adminpass12345")
    assert admin.post("/admin/messages/1/delete").get_json()["ok"]
    assert client.get("/community/history").get_json() == []
    assert client.post("/admin/messages/1/delete").status_code == 403


def test_health(client):
    assert client.get("/healthz").get_json()["status"] == "ok"
    assert client.get("/readyz").status_code == 200


# ---------------------------------------------------------------- load-handling features
def set_mode(app, mode):
    admin = app.test_client()
    login(admin, "member.admin@gmail.com", "adminpass12345")
    assert admin.post("/admin/chat-mode", data={"mode": mode}).status_code in (200, 302)
    return admin


def test_chat_mode_admins_only(app, client):
    signup(client)
    admin = set_mode(app, "admins")
    member = socketio.test_client(app, flask_test_client=client)
    r = member.emit("send_message", {"body": "hi"}, callback=True)
    assert r["ok"] is False and "announcement" in r["error"].lower()
    up = client.post("/community/upload", data={"file": (io.BytesIO(b"x"), "a.txt")}, content_type="multipart/form-data")
    assert up.status_code == 403
    boss = socketio.test_client(app, flask_test_client=admin)
    assert boss.emit("send_message", {"body": "announcement"}, callback=True)["ok"] is True


def test_slow_mode_limits_members_not_admin(app, client):
    signup(client)
    admin = set_mode(app, "slow")
    member = socketio.test_client(app, flask_test_client=client)
    assert member.emit("send_message", {"body": "one"}, callback=True)["ok"] is True
    r = member.emit("send_message", {"body": "two"}, callback=True)
    assert r["ok"] is False and "Slow mode" in r["error"]
    boss = socketio.test_client(app, flask_test_client=admin)
    assert boss.emit("send_message", {"body": "a"}, callback=True)["ok"] is True
    assert boss.emit("send_message", {"body": "b"}, callback=True)["ok"] is True


def test_chat_mode_permissions_and_validation(app, client):
    signup(client)
    assert client.post("/admin/chat-mode", data={"mode": "open"}).status_code == 403
    admin = app.test_client()
    login(admin, "member.admin@gmail.com", "adminpass12345")
    assert admin.post("/admin/chat-mode", data={"mode": "banana"}).status_code == 400


def test_new_socket_gets_current_mode(app, client):
    signup(client)
    set_mode(app, "admins")
    sio = socketio.test_client(app, flask_test_client=client)
    events = [e for e in sio.get_received() if e["name"] == "chat_mode"]
    assert events
    args = events[0]["args"]
    payload = args[0] if isinstance(args, list) else args
    assert payload["mode"] == "admins"


def test_upload_url_server_mode_and_validation(app, client):
    signup(client)
    assert client.post("/community/upload-url", json={"filename": "a.png", "size": 1000}).get_json() == {"ok": True, "mode": "server"}
    assert client.post("/community/upload-url", json={"filename": "evil.exe", "size": 10}).status_code == 400
    assert client.post("/community/upload-url", json={"filename": "a.png", "size": 10**9}).status_code == 400
    assert client.post("/community/upload-url", json={"filename": "a.png"}).status_code == 400


def test_direct_s3_upload_flow(app, client, monkeypatch):
    from app import storage

    app.config["S3_BUCKET"] = "test-bucket"
    meta = {"key": "a" * 32 + ".png", "name": "pic.png", "mime": "image/png"}
    monkeypatch.setattr(storage, "presign_upload", lambda name, ext: ({"url": "https://b.s3.amazonaws.com/", "fields": {"key": "uploads/k"}}, dict(meta)))
    monkeypatch.setattr(storage, "object_ok", lambda key, max_bytes: True)
    monkeypatch.setattr(storage, "file_url", lambda key, name=None, is_image=True: "https://cdn.example/" + key)
    signup(client)
    cfg = client.post("/community/upload-url", json={"filename": "pic.png", "size": 1234}).get_json()
    assert cfg["mode"] == "s3" and cfg["token"] and cfg["url"].startswith("https://")

    done = client.post("/community/upload-complete", json={"token": cfg["token"], "caption": "hi"})
    assert done.get_json()["ok"] is True
    last = client.get("/community/history").get_json()[-1]
    assert last["body"] == "hi" and last["file"]["is_image"] and last["file"]["url"].startswith("https://cdn.example/")

    assert client.post("/community/upload-complete", json={"token": cfg["token"]}).status_code == 409   # replay
    other = app.test_client()
    signup(other, email="member.b@gmail.com")
    assert other.post("/community/upload-complete", json={"token": cfg["token"]}).status_code == 403     # someone else's receipt
    assert client.post("/community/upload-complete", json={"token": "nonsense"}).status_code == 400


def test_s3_upload_missing_file_rejected(app, client, monkeypatch):
    from app import storage

    app.config["S3_BUCKET"] = "test-bucket"
    monkeypatch.setattr(storage, "presign_upload", lambda name, ext: ({"url": "https://b/", "fields": {}}, {"key": "b" * 32 + ".png", "name": "p.png", "mime": "image/png"}))
    monkeypatch.setattr(storage, "object_size", lambda key: None)   # browser never actually uploaded it
    signup(client)
    cfg = client.post("/community/upload-url", json={"filename": "p.png", "size": 10}).get_json()
    assert client.post("/community/upload-complete", json={"token": cfg["token"]}).status_code == 400


def test_connection_counter(app, client):
    from app import metrics

    signup(client)
    before = metrics.current()
    sio = socketio.test_client(app, flask_test_client=client)
    assert metrics.current() == before + 1
    sio.disconnect()
    assert metrics.current() == before


def test_chat_history_after_cursor(app, client):
    signup(client)
    with app.app_context():
        db.session.add_all([
            Message(user_id=1, body="one"),
            Message(user_id=1, body="two"),
            Message(user_id=1, body="three"),
        ])
        db.session.commit()
    rows = client.get("/community/history?after=1").get_json()
    assert [r["body"] for r in rows] == ["two", "three"]


def test_password_change_invalidates_old_session(app, client):
    signup(client)
    old = client.get("/community")
    assert old.status_code == 200
    with app.app_context():
        user = User.query.filter_by(email="member.a@gmail.com").one()
        user.session_version = 1
        db.session.commit()
    assert client.get("/community").status_code in (302, 401)


def test_csp_uses_nonce_and_forbids_inline_handlers(client):
    resp = client.get("/login")
    csp = resp.headers["Content-Security-Policy"]
    assert "script-src-attr 'none'" in csp
    assert "'unsafe-inline'" not in csp.split("style-src")[0]  # scripts: nonce only
    import re
    nonce = re.search(r"'nonce-([^']+)'", csp).group(1)
    assert f'nonce="{nonce}"' in resp.get_data(as_text=True)
    assert " onclick=" not in resp.get_data(as_text=True)


def test_flood_guard_limits_and_is_bounded(app):
    from app import throttle
    with app.app_context():
        results = [throttle.within_limit("flood:test", 10, 10) for _ in range(12)]
    assert results[:10] == [True] * 10 and results[10:] == [False, False]


def _make_admin(app, email="member.boss@gmail.com", pw="Password@12345"):
    with app.app_context():
        u = User(name="Boss", email=email, role="admin")
        u.set_password(pw)
        db.session.add(u)
        db.session.commit()
        return u.id


def test_like_toggles_and_counts(app, client):
    from app.models import Post
    admin_id = _make_admin(app)
    with app.app_context():
        db.session.add(Post(title="T", body="B", author_id=admin_id))
        db.session.commit()
    assert client.post("/posts/1/like").status_code in (302, 401)  # members only
    signup(client)
    r = client.post("/posts/1/like")
    assert r.get_json() == {"liked": True, "count": 1}
    r = client.post("/posts/1/like")
    assert r.get_json() == {"liked": False, "count": 0}


def test_admin_excel_export_has_member_details(app, client):
    from io import BytesIO
    from openpyxl import load_workbook
    _make_admin(app)
    signup(client, email="member.m@gmail.com", name="Meena", dob="2000-02-03", occupation="Teacher")
    client.post("/logout")
    login(client, "member.boss@gmail.com", "Password@12345")
    r = client.get("/admin/users.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.mimetype
    ws = load_workbook(BytesIO(r.data)).active
    header = [c.value for c in ws[1]]
    assert header[:6] == ["ID", "Name", "Email", "Date of birth", "Age", "Occupation"]
    rows = {row[2].value: row for row in ws.iter_rows(min_row=2)}
    assert rows["member.m@gmail.com"][1].value == "Meena" and rows["member.m@gmail.com"][5].value == "Teacher"


def test_excel_export_is_admin_only(app, client):
    signup(client)
    assert client.get("/admin/users.xlsx").status_code == 403


def test_public_feed_image_is_visible_to_visitors_but_chat_files_are_not(app, client, tmp_path):
    from app.models import Post
    admin_id = _make_admin(app)
    key = "a" * 32 + ".png"
    (tmp_path / key).write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 20)
    other = "b" * 32 + ".png"
    (tmp_path / other).write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 20)
    with app.app_context():
        db.session.add(Post(title="Pic", body="B", author_id=admin_id, file_key=key, file_name="p.png", file_mime="image/png"))
        db.session.commit()
    assert client.get(f"/files/{key}").status_code == 200
    assert client.get(f"/files/{other}").status_code in (302, 401)


def test_session_expires_after_one_hour(app, client):
    import time
    signup(client)
    assert client.get("/community").status_code == 200
    with client.session_transaction() as s:
        s["_login_at"] = int(time.time()) - app.config["SESSION_LIFETIME_SECONDS"] - 5
    r = client.get("/community")
    assert r.status_code == 302 and "/login" in r.headers["Location"] and "expired=1" in r.headers["Location"]
    assert client.get("/community").status_code == 302  # really logged out, not just redirected once


def test_expired_session_gets_json_401_for_fetch_calls(app, client):
    import time
    signup(client)
    with client.session_transaction() as s:
        s["_login_at"] = int(time.time()) - app.config["SESSION_LIFETIME_SECONDS"] - 5
    r = client.post("/posts/1/like", headers={"X-Requested-With": "fetch"})
    assert r.status_code == 401 and r.get_json()["ok"] is False


def test_login_has_no_remember_me_and_sets_one_hour_clock(app, client):
    assert app.config["SESSION_LIFETIME_SECONDS"] == 3600
    assert b"Keep me logged in" not in client.get("/login").data
    signup(client)
    with client.session_transaction() as s:
        assert s.permanent and s["_login_at"] > 0


def test_section_pages_need_login_and_show_only_their_posts(app, client):
    from app.models import Post
    admin_id = _make_admin(app)
    with app.app_context():
        db.session.add(Post(title="Home item", body="h", section="home", author_id=admin_id))
        db.session.add(Post(title="Budget news", body="n", section="news", author_id=admin_id))
        db.session.add(Post(title="Index fund pick", body="m", section="mutual_funds", author_id=admin_id))
        db.session.add(Post(title="Stock idea", body="s", section="stocks", author_id=admin_id))
        db.session.commit()
    for url in ("/news", "/mutual-funds", "/stocks"):
        assert client.get(url).status_code in (302, 401)
    signup(client)
    news = client.get("/news").get_data(as_text=True)
    assert "Budget news" in news and "Index fund pick" not in news and "Stock idea" not in news
    assert "Index fund pick" in client.get("/mutual-funds").get_data(as_text=True)
    assert "Stock idea" in client.get("/stocks").get_data(as_text=True)
    home = client.get("/").get_data(as_text=True)
    # section posts never appear as cards in the home feed (the desk preview may name the newest headline for members)
    assert "Home item" in home and home.count('class="ig-card') == 1
    visitor_home = app.test_client().get("/").get_data(as_text=True)
    assert "Budget news" not in visitor_home  # section headlines are for members only


def test_admin_publishes_to_a_section(app, client):
    from app.models import Post
    _make_admin(app)
    login(client, "member.boss@gmail.com", "Password@12345")
    r = client.post("/admin/posts/new", data={"title": "NAV update", "body": "details", "section": "mutual_funds",
                                              "category": "info"})
    assert r.status_code == 302 and r.headers["Location"].endswith("/mutual-funds")
    with app.app_context():
        assert Post.query.filter_by(title="NAV update").one().section == "mutual_funds"
    r = client.post("/admin/posts/new", data={"title": "Bad", "body": "x", "section": "hacked"})
    with app.app_context():
        assert Post.query.filter_by(title="Bad").one().section == "home"  # unknown value falls back safely


def test_after_login_and_signup_members_land_on_the_home_page(app, client):
    r = client.post("/register", data={"name": "Asha", "email": "member.n@gmail.com", "dob": "1999-05-17", "occupation": "Engineer",
                                       "password": "Password@12345", "confirm": "Password@12345"})
    assert r.status_code == 302 and r.headers["Location"].endswith("/")
    client.post("/logout")
    r = client.post("/login", data={"email": "member.n@gmail.com", "password": "Password@12345"})
    assert r.status_code == 302 and r.headers["Location"].endswith("/")


# ---------------------------------------------------------------- v16: e-mail, passwords, videos, documents
def test_email_and_password_rules(app, client):
    from app.validators import check_email, check_password
    with app.app_context():
        assert check_email("Asha.Kumar@Gmail.com")[1] is None
        assert check_email("asha@fakedomain.xyz")[1]
        assert check_email("ab@gmail.com")[1]           # Gmail names need 6+ characters
        assert check_email("not-an-email")[1]
        assert check_password("Password@12345") is None
        assert check_password("password12345")           # no special character
        assert check_password("Password@abcdef")        # no number
        assert check_password("123456@7890123")         # no letter
        assert check_password("Pa@1")                   # too short
    signup(client, email="x@fakedomain.xyz")
    with app.app_context():
        assert User.query.filter_by(email="x@fakedomain.xyz").first() is None


def test_forgot_password_sends_mail(app, client, monkeypatch):
    sent = []
    import app.auth as auth_mod
    monkeypatch.setattr(auth_mod, "send_mail", lambda to, subj, text, html=None: sent.append((to, text)) or True)
    signup(client, email="reset.user@gmail.com")
    client.post("/logout")
    client.post("/forgot-password", data={"email": "reset.user@gmail.com"})
    assert sent and sent[0][0] == "reset.user@gmail.com" and "/reset-password/" in sent[0][1]


def test_video_limit_is_50mb_and_others_15mb(app):
    from app import storage
    with app.app_context():
        assert storage.size_error("mp4", 40 * 1024 * 1024) is None
        assert storage.size_error("mp4", 51 * 1024 * 1024)
        assert storage.size_error("pdf", 14 * 1024 * 1024) is None
        assert storage.size_error("pdf", 16 * 1024 * 1024)


def test_documents_page_lists_uploads(app, client):
    from app.models import Message
    signup(client, email="doc.member@gmail.com")
    with app.app_context():
        u = User.query.filter_by(email="doc.member@gmail.com").one()
        db.session.add(Message(user_id=u.id, body="report", file_key="a" * 32 + ".pdf", file_name="report.pdf",
                               file_mime="application/pdf", channel="investors"))
        db.session.commit()
    html = client.get("/documents").get_data(as_text=True)
    assert "report.pdf" in html and "Investors" in html
