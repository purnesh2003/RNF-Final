"""Admin area: dashboard, user management, quizzes, polls, opinions, moderation."""
import csv
import io
from datetime import date, timezone
import json
import secrets

from flask import Blueprint, Response, abort, current_app, flash, jsonify, redirect, render_template, request, send_file, url_for
from flask_login import current_user
from sqlalchemy import func

from . import analytics, chatmode, gemini_quiz, notify, quiz_notify, storage
from .extensions import db, limiter, socketio
from .utils import is_safe_redirect
from .models import Message, Poll, PollOption, PollVote, Post, PostLike, Question, Quiz, QuizAttempt, User
from .utils import admin_required

bp = Blueprint("admin", __name__, url_prefix="/admin")


@bp.route("/")
@admin_required
def dashboard():
    stats = {
        "users": User.query.count(),
        "active": User.query.filter_by(active=True).count(),
        "messages": Message.query.filter_by(deleted=False).count(),
        "quizzes": Quiz.query.count(),
        "polls": Poll.query.count(),
        "posts": Post.query.count(),
    }
    recent = User.query.order_by(User.created_at.desc()).limit(8).all()
    return render_template("admin/dashboard.html", stats=stats, recent=recent, chat_mode=chatmode.get_mode())


@bp.route("/analytics")
@admin_required
def analytics_page():
    """Graphs of members, activity, occupations, quiz attendance and chat (read-only)."""
    days = request.args.get("days", type=int) or analytics.DEFAULT_RANGE
    return render_template("admin/analytics.html", data=analytics.build(days))


@bp.route("/chat-mode", methods=["POST"])
@admin_required
def chat_mode():
    """open = everyone can post | slow = one message per member every N seconds | admins = announcement-only."""
    mode = request.form.get("mode") or (request.get_json(silent=True) or {}).get("mode")
    if mode not in chatmode.MODES:
        return {"ok": False, "error": "Unknown mode."}, 400
    chatmode.set_mode(mode)
    socketio.emit("chat_mode", {"mode": mode}, to="community")
    if request.headers.get("X-Requested-With") == "fetch":
        return {"ok": True, "mode": mode}
    flash(f"Chat mode set to '{mode}'.", "success")
    return redirect(url_for("admin.dashboard"))


# ---------------------------------------------------------------- users
@bp.route("/users")
@admin_required
def users():
    q = request.args.get("q", "").strip()
    query = User.query
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(User.name.ilike(like), User.email.ilike(like), User.occupation.ilike(like)))
    page = request.args.get("page", 1, type=int)
    users = query.order_by(User.created_at.desc()).paginate(page=page, per_page=50, error_out=False)
    return render_template("admin/users.html", users=users, q=q)


def _csv_safe(v):
    v = "" if v is None else str(v)
    return "'" + v if v[:1] in ("=", "+", "-", "@") else v  # block spreadsheet formula injection


@bp.route("/users.csv")
@admin_required
def users_csv():
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["id", "name", "email", "dob", "occupation", "role", "active", "muted", "joined", "last_login"])
    for u in User.query.order_by(User.id):
        w.writerow([u.id, _csv_safe(u.name), _csv_safe(u.email), u.dob or "", _csv_safe(u.occupation), u.role, u.active, u.muted,
                    u.created_at, u.last_login or ""])
    return Response(buf.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": "attachment; filename=users.csv"})


def _naive_utc(dt):
    """Excel cannot store time zones; write UTC wall-clock time."""
    if dt is None:
        return None
    return dt.astimezone(timezone.utc).replace(tzinfo=None) if dt.tzinfo else dt


def _age(dob):
    if not dob:
        return None
    today = date.today()
    return today.year - dob.year - ((today.month, today.day) < (dob.month, dob.day))


@bp.route("/users.xlsx")
@admin_required
def users_xlsx():
    """Download every member's details as a formatted Excel workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Members"
    headers = ["ID", "Name", "Email", "Date of birth", "Age", "Occupation", "Role", "Status", "Muted",
               "Joined (UTC)", "Last login (UTC)"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F172A")
        c.alignment = Alignment(vertical="center")
    text_cols = {2, 3, 6}  # free text typed by members: store as text so "=..." is never run as a formula
    for u in User.query.order_by(User.id):
        ws.append([u.id, u.name, u.email, u.dob, _age(u.dob), u.occupation or "", u.role,
                   "Active" if u.active else "Suspended", "Yes" if u.muted else "No",
                   _naive_utc(u.created_at), _naive_utc(u.last_login)])
        row = ws.max_row
        for col in text_cols:
            cell = ws.cell(row=row, column=col)
            if isinstance(cell.value, str):
                cell.data_type = "s"
        ws.cell(row=row, column=4).number_format = "DD-MMM-YYYY"
        ws.cell(row=row, column=10).number_format = "DD-MMM-YYYY HH:MM"
        ws.cell(row=row, column=11).number_format = "DD-MMM-YYYY HH:MM"
    widths = [7, 26, 34, 15, 7, 24, 9, 11, 8, 20, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    current_app.logger.info("admin %s exported the member list (%d rows)", current_user.id, ws.max_row - 1)
    return send_file(buf, as_attachment=True, download_name=f"members-{date.today().isoformat()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/users/<int:user_id>/<action>", methods=["POST"])
@admin_required
def user_action(user_id, action):
    user = db.get_or_404(User, user_id)
    is_self = user.id == current_user.id
    if is_self and action in ("suspend", "delete", "demote"):
        flash("You can't do that to your own account.", "warning")
        return redirect(request.referrer or url_for("admin.users"))

    if action == "suspend":
        user.active = False
    elif action == "activate":
        user.active = True
    elif action == "mute":
        user.muted = True
    elif action == "unmute":
        user.muted = False
    elif action == "promote":
        user.role = "admin"
    elif action == "demote":
        user.role = "user"
    elif action == "reset-password":
        temp = secrets.token_urlsafe(9)
        user.set_password(temp)
        user.session_version += 1
        user.must_change_password = True
        flash(f"Temporary password for {user.email}: {temp}  (shown once — share it privately; "
              f"they must change it at next login)", "warning")
    elif action == "delete":
        if Post.query.filter_by(author_id=user.id).first():
            flash("This user authored opinions; demote/suspend instead of deleting.", "warning")
            return redirect(request.referrer or url_for("admin.users"))
        Message.query.filter_by(user_id=user.id).delete()
        PollVote.query.filter_by(user_id=user.id).delete()
        PostLike.query.filter_by(user_id=user.id).delete()
        QuizAttempt.query.filter_by(user_id=user.id).delete()
        db.session.delete(user)
    else:
        abort(404)
    db.session.commit()
    if action != "reset-password":
        flash(f"{action.replace('-', ' ').title()} done for {user.email}.", "success")
    return redirect(request.referrer or url_for("admin.users"))


# ---------------------------------------------------------------- moderation
@bp.route("/messages/<int:msg_id>/delete", methods=["POST"])
@admin_required
def delete_message(msg_id):
    msg = db.get_or_404(Message, msg_id)
    msg.deleted = True
    db.session.commit()
    socketio.emit("message_deleted", {"id": msg.id}, to="community")
    return {"ok": True}


# ---------------------------------------------------------------- opinions
def _post_form(selected):
    """Template values for the 'new post' form (selected = which place is pre-chosen)."""
    return {"categories": Post.CATEGORIES, "sections": Post.SECTIONS,
            "selected": selected if selected in Post.SECTIONS else "home"}


@bp.route("/posts/new", methods=["GET", "POST"])
@admin_required
def new_post():
    if request.method == "POST":
        section = request.form.get("section", "home")
        if section not in Post.SECTIONS:
            section = "home"
        title, body = request.form.get("title", "").strip()[:200], request.form.get("body", "").strip()
        if not title or not body:
            flash("Title and text are required.", "danger")
            return render_template("admin/new_post.html", **_post_form(section)), 400
        category = request.form.get("category", "news")
        if category not in Post.CATEGORIES:
            category = "news"
        post = Post(title=title, body=body, author_id=current_user.id, category=category, section=section,
                    pinned=request.form.get("pinned") == "1")
        file = request.files.get("file")
        if file and file.filename:
            ok, err, ext = storage.validate(file)
            if not ok:
                flash(err, "danger")
                return render_template("admin/new_post.html", **_post_form(section)), 400
            too_big = storage.size_error(ext, storage.stream_size(file))
            if too_big:
                flash(too_big, "danger")
                return render_template("admin/new_post.html", **_post_form(section)), 400
            saved = storage.save_upload(file, ext)
            post.file_key, post.file_name, post.file_mime = saved["key"], saved["name"], saved["mime"]
        db.session.add(post)
        db.session.commit()
        notify.announce_post(post)  # e-mails members; never blocks or fails the page
        flash(f"Published to {Post.SECTIONS[section].split(' ', 1)[1]}.", "success")
        from .main import SECTION_ENDPOINTS
        return redirect(url_for(SECTION_ENDPOINTS[section]))
    return render_template("admin/new_post.html", **_post_form(request.args.get("section", "home")))


@bp.route("/posts/<int:post_id>/delete", methods=["POST"])
@admin_required
def delete_post(post_id):
    db.session.delete(db.get_or_404(Post, post_id))
    db.session.commit()
    flash("Post deleted.", "success")
    return redirect(request.form.get("next") if is_safe_redirect(request.form.get("next")) else url_for("main.home"))


# ---------------------------------------------------------------- quizzes
def _parse_questions(raw):
    try:
        items = json.loads(raw)
    except (TypeError, ValueError):
        return None, "Could not read the questions."
    if not isinstance(items, list) or not items:
        return None, "Add at least one question."
    out = []
    for i, it in enumerate(items[:100], 1):
        if not isinstance(it, dict) or not isinstance(it.get("options"), list):
            return None, f"Question {i} is malformed."
        text = str(it.get("text", "")).strip()[:500]
        original = [str(o).strip()[:200] for o in it["options"]]
        opts = [o for o in original if o]
        correct = it.get("correct")
        explanation = str(it.get("explanation") or "").strip()[:600]
        if not text:
            return None, f"Question {i} has no text."
        if len(opts) < 2:
            return None, f"Question {i} needs at least 2 options."
        if not isinstance(correct, int) or isinstance(correct, bool) or not 0 <= correct < len(original):
            return None, f"Question {i} needs a correct answer."
        if not original[correct]:
            return None, f"Question {i}: the correct answer can't be blank."
        # blank options are dropped, so shift the correct index accordingly
        new_correct = sum(1 for o in original[:correct] if o)
        out.append((text, opts, new_correct, explanation))
    return out, None


def _ai_ready():
    return bool(current_app.config.get("GEMINI_API_KEY"))


@bp.route("/quizzes/new", methods=["GET", "POST"])
@admin_required
def new_quiz():
    if request.method == "POST":
        title = request.form.get("title", "").strip()[:200]
        parsed, err = _parse_questions(request.form.get("questions_json"))
        if not title:
            err = "Give the quiz a title."
        if err:
            flash(err, "danger")
            return render_template("admin/new_quiz.html", ai_ready=_ai_ready(), max_q=current_app.config["QUIZ_MAX_QUESTIONS"]), 400
        quiz = Quiz(title=title, description=request.form.get("description", "").strip())
        if request.form.get("source") == "ai":  # quiz drafted by Gemini and reviewed by the admin
            quiz.source = "ai"
            quiz.topic = request.form.get("topic", "").strip()[:200] or None
            diff = request.form.get("difficulty", "").strip().lower()
            quiz.difficulty = diff if diff in gemini_quiz.DIFFICULTIES else None
        quiz.questions = [Question(text=t, options=o, correct_index=c, explanation=x or None)
                          for t, o, c, x in parsed]
        db.session.add(quiz)
        db.session.commit()
        started = quiz_notify.announce_quiz(quiz.id)  # e-mails members once; never raises into the page
        flash("Quiz published." + (" Members are being notified by e-mail." if started else ""), "success")
        return redirect(url_for("main.quizzes"))
    return render_template("admin/new_quiz.html", ai_ready=_ai_ready(), max_q=current_app.config["QUIZ_MAX_QUESTIONS"])


@bp.route("/quizzes/generate", methods=["POST"])
@admin_required
@limiter.limit("20 per hour")
def generate_quiz():
    """Ask Gemini for a draft. Returns JSON only; nothing is saved here, so a failure can never leave
    a half-made quiz behind. The admin reviews/edits the draft and then publishes it via new_quiz()."""
    data = request.get_json(silent=True) or request.form
    try:
        number = int(data.get("count"))
    except (TypeError, ValueError):
        number = 0
    try:
        cfg = current_app.config
        result = gemini_quiz.generate_quiz(
            data.get("topic", ""), number, data.get("difficulty", ""),
            api_key=cfg["GEMINI_API_KEY"], model=cfg["GEMINI_MODEL"],
            timeout=cfg["GEMINI_TIMEOUT"], max_questions=cfg["QUIZ_MAX_QUESTIONS"])
    except gemini_quiz.QuizGenerationError as e:
        return jsonify(ok=False, error=str(e)), 502 if "Gemini" in str(e) else 400
    except Exception:  # anything unexpected: log it (no secrets) and stay friendly
        current_app.logger.exception("Quiz generation crashed")
        return jsonify(ok=False, error="Something went wrong while generating the quiz. Please try again."), 500
    return jsonify(ok=True, **result)


@bp.route("/quizzes/<int:quiz_id>/<action>", methods=["POST"])
@admin_required
def quiz_action(quiz_id, action):
    quiz = db.get_or_404(Quiz, quiz_id)
    if action == "toggle":
        quiz.is_open = not quiz.is_open
    elif action == "delete":
        QuizAttempt.query.filter_by(quiz_id=quiz.id).delete()
        db.session.delete(quiz)
    else:
        abort(404)
    db.session.commit()
    return redirect(url_for("main.quizzes"))


@bp.route("/quizzes/<int:quiz_id>/leaderboard")
@admin_required
def leaderboard(quiz_id):
    quiz = db.get_or_404(Quiz, quiz_id)
    rows = QuizAttempt.query.filter_by(quiz_id=quiz.id).order_by(
        QuizAttempt.score.desc(), QuizAttempt.created_at).all()
    return render_template("admin/leaderboard.html", quiz=quiz, rows=rows)


# ---------------------------------------------------------------- polls
@bp.route("/polls/new", methods=["GET", "POST"])
@admin_required
def new_poll():
    if request.method == "POST":
        question = request.form.get("question", "").strip()[:300]
        options = [o.strip()[:200] for o in request.form.get("options", "").splitlines() if o.strip()]
        if not question or len(options) < 2:
            flash("A poll needs a question and at least 2 options (one per line).", "danger")
            return render_template("admin/new_poll.html"), 400
        poll = Poll(question=question)
        poll.options = [PollOption(text=o) for o in options[:20]]
        db.session.add(poll)
        db.session.commit()
        flash("Poll published.", "success")
        return redirect(url_for("main.polls"))
    return render_template("admin/new_poll.html")


@bp.route("/polls/<int:poll_id>/<action>", methods=["POST"])
@admin_required
def poll_action(poll_id, action):
    poll = db.get_or_404(Poll, poll_id)
    if action == "toggle":
        poll.is_open = not poll.is_open
    elif action == "delete":
        PollVote.query.filter_by(poll_id=poll.id).delete()
        db.session.delete(poll)
    else:
        abort(404)
    db.session.commit()
    return redirect(url_for("main.polls"))
